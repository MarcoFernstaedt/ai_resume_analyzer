"""Resume export service - generates Word documents from resume data."""
import io
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _set_cell_background(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def export_harvard_template(resume_data: dict) -> bytes:
    """Generate a Harvard-style resume Word document."""
    doc = Document()

    # Page margins
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    # Name
    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_para.add_run(resume_data.get('name', 'Your Name'))
    name_run.bold = True
    name_run.font.size = Pt(18)

    # Contact line
    contact_parts = [
        resume_data.get('email', ''),
        resume_data.get('phone', ''),
        resume_data.get('location', ''),
        resume_data.get('linkedin', ''),
        resume_data.get('github', ''),
    ]
    contact_line = ' | '.join(p for p in contact_parts if p)
    if contact_line:
        contact_para = doc.add_paragraph(contact_line)
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_para.runs[0].font.size = Pt(10)

    doc.add_paragraph()

    def add_section_header(title):
        para = doc.add_paragraph()
        run = para.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(11)
        para.paragraph_format.space_before = Pt(6)
        # Add bottom border
        p_pr = para._p.get_or_add_pPr()
        p_bdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '000000')
        p_bdr.append(bottom)
        p_pr.append(p_bdr)

    # Summary
    if resume_data.get('summary'):
        add_section_header('Professional Summary')
        doc.add_paragraph(resume_data['summary'])

    # Experience
    if resume_data.get('experience'):
        add_section_header('Work Experience')
        for exp in resume_data['experience']:
            p = doc.add_paragraph()
            r = p.add_run(exp.get('title', '') + ' — ' + exp.get('company', ''))
            r.bold = True
            r.font.size = Pt(11)
            p2 = doc.add_paragraph()
            p2.add_run(exp.get('dates', '') + ' | ' + exp.get('location', '')).italic = True
            for bullet in exp.get('bullets', []):
                doc.add_paragraph(bullet, style='List Bullet')

    # Education
    if resume_data.get('education'):
        add_section_header('Education')
        for edu in resume_data['education']:
            p = doc.add_paragraph()
            r = p.add_run(edu.get('degree', '') + ' — ' + edu.get('school', ''))
            r.bold = True
            doc.add_paragraph(edu.get('dates', '') + ' | GPA: ' + edu.get('gpa', 'N/A'))

    # Skills
    if resume_data.get('skills'):
        add_section_header('Technical Skills')
        skills_text = ' | '.join(resume_data['skills']) if isinstance(resume_data['skills'], list) else resume_data['skills']
        doc.add_paragraph(skills_text)

    # Projects
    if resume_data.get('projects'):
        add_section_header('Projects')
        for proj in resume_data['projects']:
            p = doc.add_paragraph()
            r = p.add_run(proj.get('name', ''))
            r.bold = True
            if proj.get('tech'):
                p.add_run(' | ' + proj['tech'])
            doc.add_paragraph(proj.get('description', ''))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_google_template(resume_data: dict) -> bytes:
    """Generate a Google-style (clean, minimal) resume Word document."""
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    # Header with accent color
    name_para = doc.add_paragraph()
    name_run = name_para.add_run(resume_data.get('name', 'Your Name'))
    name_run.bold = True
    name_run.font.size = Pt(20)
    name_run.font.color.rgb = RGBColor(0x42, 0x85, 0xF4)  # Google blue

    contact_parts = [
        resume_data.get('email', ''),
        resume_data.get('phone', ''),
        resume_data.get('linkedin', ''),
        resume_data.get('github', ''),
    ]
    contact_line = '  •  '.join(p for p in contact_parts if p)
    if contact_line:
        doc.add_paragraph(contact_line).runs[0].font.size = Pt(9)

    def add_section(title):
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(10)
        run = para.add_run(title)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x42, 0x85, 0xF4)

    if resume_data.get('summary'):
        add_section('Summary')
        doc.add_paragraph(resume_data['summary']).runs[0].font.size = Pt(10)

    if resume_data.get('experience'):
        add_section('Experience')
        for exp in resume_data['experience']:
            p = doc.add_paragraph()
            r = p.add_run(exp.get('title', ''))
            r.bold = True
            r.font.size = Pt(11)
            p.add_run(f"  |  {exp.get('company', '')}  |  {exp.get('dates', '')}")
            for bullet in exp.get('bullets', []):
                bp = doc.add_paragraph(style='List Bullet')
                bp.add_run(bullet).font.size = Pt(10)

    if resume_data.get('skills'):
        add_section('Technical Skills')
        skills_text = ', '.join(resume_data['skills']) if isinstance(resume_data['skills'], list) else resume_data['skills']
        doc.add_paragraph(skills_text).runs[0].font.size = Pt(10)

    if resume_data.get('education'):
        add_section('Education')
        for edu in resume_data['education']:
            p = doc.add_paragraph()
            p.add_run(edu.get('degree', '')).bold = True
            doc.add_paragraph(f"{edu.get('school', '')} | {edu.get('dates', '')}")

    if resume_data.get('projects'):
        add_section('Projects')
        for proj in resume_data['projects']:
            p = doc.add_paragraph()
            p.add_run(proj.get('name', '')).bold = True
            if proj.get('url'):
                p.add_run(f"  —  {proj['url']}")
            doc.add_paragraph(proj.get('description', '')).runs[0].font.size = Pt(10)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_plain_text(resume_data: dict) -> str:
    """Generate plain text resume (ATS-friendly)."""
    lines = []

    lines.append(resume_data.get('name', 'Your Name').upper())
    lines.append('=' * 60)

    contact = ' | '.join(filter(None, [
        resume_data.get('email', ''),
        resume_data.get('phone', ''),
        resume_data.get('location', ''),
        resume_data.get('linkedin', ''),
        resume_data.get('github', ''),
    ]))
    if contact:
        lines.append(contact)
    lines.append('')

    if resume_data.get('summary'):
        lines.append('PROFESSIONAL SUMMARY')
        lines.append('-' * 40)
        lines.append(resume_data['summary'])
        lines.append('')

    if resume_data.get('skills'):
        lines.append('TECHNICAL SKILLS')
        lines.append('-' * 40)
        skills = resume_data['skills']
        lines.append(', '.join(skills) if isinstance(skills, list) else skills)
        lines.append('')

    if resume_data.get('experience'):
        lines.append('WORK EXPERIENCE')
        lines.append('-' * 40)
        for exp in resume_data['experience']:
            lines.append(f"{exp.get('title', '')} | {exp.get('company', '')} | {exp.get('dates', '')}")
            for bullet in exp.get('bullets', []):
                lines.append(f"  • {bullet}")
            lines.append('')

    if resume_data.get('education'):
        lines.append('EDUCATION')
        lines.append('-' * 40)
        for edu in resume_data['education']:
            lines.append(f"{edu.get('degree', '')} | {edu.get('school', '')} | {edu.get('dates', '')}")
        lines.append('')

    if resume_data.get('projects'):
        lines.append('PROJECTS')
        lines.append('-' * 40)
        for proj in resume_data['projects']:
            lines.append(f"{proj.get('name', '')} | {proj.get('tech', '')}")
            lines.append(proj.get('description', ''))
            lines.append('')

    return '\n'.join(lines)
