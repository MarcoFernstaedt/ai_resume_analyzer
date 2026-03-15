"""Resume file parsing service - supports PDF and DOCX."""
import io
import logging

logger = logging.getLogger(__name__)


def parse_resume(file_obj) -> str:
    """Extract plain text from a resume file (PDF or DOCX)."""
    filename = getattr(file_obj, 'name', '').lower()

    if filename.endswith('.pdf'):
        return _parse_pdf(file_obj)
    elif filename.endswith('.docx'):
        return _parse_docx(file_obj)
    elif filename.endswith('.txt'):
        return file_obj.read().decode('utf-8', errors='replace')
    else:
        raise ValueError(f'Unsupported file type: {filename}')


def _parse_pdf(file_obj) -> str:
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(file_obj) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
        return '\n'.join(text_parts)
    except Exception as e:
        logger.error(f'PDF parsing error: {e}')
        raise ValueError(f'Could not parse PDF: {e}')


def _parse_docx(file_obj) -> str:
    try:
        from docx import Document
        doc = Document(file_obj)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return '\n'.join(paragraphs)
    except Exception as e:
        logger.error(f'DOCX parsing error: {e}')
        raise ValueError(f'Could not parse DOCX: {e}')
